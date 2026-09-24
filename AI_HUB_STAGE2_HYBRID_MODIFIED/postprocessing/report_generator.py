import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from postprocessing.log_reader import LogReader
from postprocessing.metrics import MetricsCalculator


class ReportGenerator:

    def __init__(self, report_directory):
        self.report_directory = Path(report_directory)

        self.report_directory.mkdir(
            parents=True,
            exist_ok=True
        )

    # =========================================================
    # DETERMINE TP / TN / FP / FN FOR ONE VIDEO
    # =========================================================

    def classification_flags(self, result):

        actual = str(
            result["ground_truth"]
        ).lower()

        predicted = str(
            result["prediction"]
        ).lower()

        tp = 0
        tn = 0
        fp = 0
        fn = 0

        if (
            actual == "accident"
            and predicted == "accident"
        ):
            tp = 1

        elif (
            actual == "normal"
            and predicted == "normal"
        ):
            tn = 1

        elif (
            actual == "normal"
            and predicted == "accident"
        ):
            fp = 1

        elif (
            actual == "accident"
            and predicted == "normal"
        ):
            fn = 1

        return tp, tn, fp, fn

    # =========================================================
    # SET COLUMN WIDTHS
    # =========================================================

    def set_widths(
        self,
        sheet,
        widths
    ):

        for column_letter, width in widths.items():

            sheet.column_dimensions[
                column_letter
            ].width = width

    # =========================================================
    # SHEET 1 — VIDEO RESULTS
    # =========================================================

    def generate_video_results_sheet(
        self,
        workbook,
        results
    ):

        sheet = workbook.active

        sheet.title = "Video Results"

        headers = [
            "Video",
            "Ground Truth",
            "Prediction",
            "TP",
            "TN",
            "FP",
            "FN",
            "First Accident Timestamp",
            "Frames Processed",
            "Detection Count",
            "Vehicle Detection Count",
            "Suspicious Frame Count",
            "Processing Time (seconds)"
        ]

        sheet.append(headers)

        # -----------------------------------------------------
        # HEADER FORMATTING
        # -----------------------------------------------------

        for cell in sheet[1]:

            cell.font = Font(
                bold=True
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )

        # -----------------------------------------------------
        # VIDEO RESULTS
        # -----------------------------------------------------

        total_tp = 0
        total_tn = 0
        total_fp = 0
        total_fn = 0

        for result in results:

            tp, tn, fp, fn = (
                self.classification_flags(
                    result
                )
            )

            total_tp += tp
            total_tn += tn
            total_fp += fp
            total_fn += fn

            sheet.append([
                result.get(
                    "video_name"
                ),

                result.get(
                    "ground_truth"
                ),

                result.get(
                    "prediction"
                ),

                tp,
                tn,
                fp,
                fn,

                result.get(
                    "first_accident_timestamp"
                ),

                result.get(
                    "frames_processed"
                ),

                result.get(
                    "detection_count"
                ),

                result.get(
                    "vehicle_detection_count"
                ),

                result.get(
                    "suspicious_frame_count"
                ),

                result.get(
                    "processing_time_seconds"
                )
            ])

        # -----------------------------------------------------
        # TOTAL ROW
        # -----------------------------------------------------

        total_row = (
            sheet.max_row + 2
        )

        sheet.cell(
            total_row,
            1,
            "TOTAL"
        )

        sheet.cell(
            total_row,
            4,
            total_tp
        )

        sheet.cell(
            total_row,
            5,
            total_tn
        )

        sheet.cell(
            total_row,
            6,
            total_fp
        )

        sheet.cell(
            total_row,
            7,
            total_fn
        )

        for cell in sheet[total_row]:

            cell.font = Font(
                bold=True
            )

        # -----------------------------------------------------
        # COLUMN WIDTHS
        # -----------------------------------------------------

        self.set_widths(
            sheet,
            {
                "A": 18,
                "B": 16,
                "C": 16,
                "D": 8,
                "E": 8,
                "F": 8,
                "G": 8,
                "H": 25,
                "I": 18,
                "J": 18,
                "K": 25,
                "L": 25,
                "M": 25
            }
        )

        # -----------------------------------------------------
        # ALIGNMENT
        # -----------------------------------------------------

        for row in sheet.iter_rows():

            for cell in row:

                cell.alignment = Alignment(
                    vertical="center",
                    wrap_text=True
                )

    # =========================================================
    # SHEET 2 — COMPLETE CALCULATIONS
    # =========================================================

    def generate_calculation_sheet(
        self,
        workbook,
        metrics
    ):

        sheet = workbook.create_sheet(
            "Complete Calculations"
        )

        tp = metrics["TP"]
        tn = metrics["TN"]
        fp = metrics["FP"]
        fn = metrics["FN"]

        # =====================================================
        # TITLE
        # =====================================================

        sheet["A1"] = (
            "COMPLETE METRIC CALCULATIONS"
        )

        sheet["A1"].font = Font(
            bold=True,
            size=14
        )

        # =====================================================
        # CONFUSION MATRIX TOTALS
        # =====================================================

        sheet["A3"] = (
            "CONFUSION MATRIX TOTALS"
        )

        sheet["A3"].font = Font(
            bold=True
        )

        sheet["A4"] = (
            "True Positive (TP)"
        )

        sheet["B4"] = tp

        sheet["A5"] = (
            "True Negative (TN)"
        )

        sheet["B5"] = tn

        sheet["A6"] = (
            "False Positive (FP)"
        )

        sheet["B6"] = fp

        sheet["A7"] = (
            "False Negative (FN)"
        )

        sheet["B7"] = fn

        total_videos = (
            tp
            + tn
            + fp
            + fn
        )

        sheet["A8"] = (
            "Total Videos"
        )

        sheet["B8"] = total_videos

        # =====================================================
        # CALCULATION TABLE HEADERS
        # =====================================================

        sheet["A10"] = "Metric"
        sheet["B10"] = "Formula"
        sheet["C10"] = "Actual Calculation"
        sheet["D10"] = "Result"

        for cell in sheet[10]:

            cell.font = Font(
                bold=True
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )

        # =====================================================
        # ACCURACY
        # =====================================================

        accuracy_denominator = (
            tp
            + tn
            + fp
            + fn
        )

        if accuracy_denominator > 0:

            accuracy = (
                (tp + tn)
                / accuracy_denominator
            )

        else:

            accuracy = 0.0

        sheet["A11"] = "Accuracy"

        sheet["B11"] = (
            "(TP + TN) / "
            "(TP + TN + FP + FN)"
        )

        sheet["C11"] = (
            f"({tp} + {tn}) / "
            f"({tp} + {tn} + {fp} + {fn})"
        )

        sheet["D11"] = (
            f"{accuracy * 100:.2f}%"
        )

        # =====================================================
        # PRECISION
        # =====================================================

        precision_denominator = (
            tp + fp
        )

        if precision_denominator > 0:

            precision = (
                tp
                / precision_denominator
            )

        else:

            precision = 0.0

        sheet["A12"] = "Precision"

        sheet["B12"] = (
            "TP / (TP + FP)"
        )

        sheet["C12"] = (
            f"{tp} / ({tp} + {fp})"
        )

        sheet["D12"] = (
            f"{precision * 100:.2f}%"
        )

        # =====================================================
        # RECALL
        # =====================================================

        recall_denominator = (
            tp + fn
        )

        if recall_denominator > 0:

            recall = (
                tp
                / recall_denominator
            )

        else:

            recall = 0.0

        sheet["A13"] = "Recall"

        sheet["B13"] = (
            "TP / (TP + FN)"
        )

        sheet["C13"] = (
            f"{tp} / ({tp} + {fn})"
        )

        sheet["D13"] = (
            f"{recall * 100:.2f}%"
        )

        # =====================================================
        # F1 SCORE
        # =====================================================

        if (
            precision + recall
        ) > 0:

            f1 = (
                2
                * precision
                * recall
                / (
                    precision
                    + recall
                )
            )

        else:

            f1 = 0.0

        sheet["A14"] = "F1 Score"

        sheet["B14"] = (
            "2 × Precision × Recall / "
            "(Precision + Recall)"
        )

        sheet["C14"] = (
            f"2 × {precision:.4f} × "
            f"{recall:.4f} / "
            f"({precision:.4f} + "
            f"{recall:.4f})"
        )

        sheet["D14"] = (
            f"{f1 * 100:.2f}%"
        )

        # =====================================================
        # SPECIFICITY
        # =====================================================

        specificity_denominator = (
            tn + fp
        )

        if specificity_denominator > 0:

            specificity = (
                tn
                / specificity_denominator
            )

        else:

            specificity = 0.0

        sheet["A15"] = "Specificity"

        sheet["B15"] = (
            "TN / (TN + FP)"
        )

        sheet["C15"] = (
            f"{tn} / ({tn} + {fp})"
        )

        sheet["D15"] = (
            f"{specificity * 100:.2f}%"
        )

        # =====================================================
        # FINAL SUMMARY
        # =====================================================

        sheet["A18"] = (
            "FINAL METRICS SUMMARY"
        )

        sheet["A18"].font = Font(
            bold=True
        )

        sheet["A19"] = "Metric"
        sheet["B19"] = "Value"

        sheet["A19"].font = Font(
            bold=True
        )

        sheet["B19"].font = Font(
            bold=True
        )

        summary = [
            ("Accuracy", accuracy),
            ("Precision", precision),
            ("Recall", recall),
            ("F1 Score", f1),
            ("Specificity", specificity)
        ]

        row = 20

        for metric_name, value in summary:

            sheet.cell(
                row,
                1,
                metric_name
            )

            sheet.cell(
                row,
                2,
                f"{value * 100:.2f}%"
            )

            row += 1

        # =====================================================
        # FORMATTING
        # =====================================================

        for row_cells in sheet.iter_rows():

            for cell in row_cells:

                cell.alignment = Alignment(
                    vertical="center",
                    wrap_text=True
                )

        self.set_widths(
            sheet,
            {
                "A": 25,
                "B": 45,
                "C": 60,
                "D": 18
            }
        )

    # =========================================================
    # GENERATE EXCEL REPORT
    # =========================================================

    def generate_excel(
        self,
        results,
        metrics,
        filename="complete_results.xlsx"
    ):

        output_path = (
            self.report_directory
            / filename
        )

        workbook = Workbook()

        # -----------------------------------------------------
        # VIDEO RESULTS
        # -----------------------------------------------------

        self.generate_video_results_sheet(
            workbook,
            results
        )

        # -----------------------------------------------------
        # COMPLETE CALCULATIONS
        # -----------------------------------------------------

        self.generate_calculation_sheet(
            workbook,
            metrics
        )

        # -----------------------------------------------------
        # SAVE
        # -----------------------------------------------------

        workbook.save(
            output_path
        )

        return output_path

    # =========================================================
    # GENERATE CSV REPORT
    # =========================================================

    def generate_csv(
        self,
        results,
        filename="video_results.csv"
    ):

        output_path = (
            self.report_directory
            / filename
        )

        fieldnames = [
            "Video",
            "Ground Truth",
            "Prediction",
            "TP",
            "TN",
            "FP",
            "FN",
            "First Accident Timestamp",
            "Frames Processed",
            "Detection Count",
            "Vehicle Detection Count",
            "Suspicious Frame Count",
            "Processing Time (seconds)"
        ]

        with open(
            output_path,
            "w",
            newline="",
            encoding="utf-8"
        ) as file:

            writer = csv.DictWriter(
                file,
                fieldnames=fieldnames
            )

            writer.writeheader()

            for result in results:

                tp, tn, fp, fn = (
                    self.classification_flags(
                        result
                    )
                )

                writer.writerow({
                    "Video":
                        result.get(
                            "video_name"
                        ),

                    "Ground Truth":
                        result.get(
                            "ground_truth"
                        ),

                    "Prediction":
                        result.get(
                            "prediction"
                        ),

                    "TP":
                        tp,

                    "TN":
                        tn,

                    "FP":
                        fp,

                    "FN":
                        fn,

                    "First Accident Timestamp":
                        result.get(
                            "first_accident_timestamp"
                        ),

                    "Frames Processed":
                        result.get(
                            "frames_processed"
                        ),

                    "Detection Count":
                        result.get(
                            "detection_count"
                        ),

                    "Vehicle Detection Count":
                        result.get(
                            "vehicle_detection_count"
                        ),

                    "Suspicious Frame Count":
                        result.get(
                            "suspicious_frame_count"
                        ),

                    "Processing Time (seconds)":
                        result.get(
                            "processing_time_seconds"
                        )
                })

        return output_path