import json
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment


# ==================================================
# PATHS
# ==================================================

PROJECT_ROOT = Path(__file__).resolve().parent

LOG_DIRECTORY = PROJECT_ROOT / "logs" / "v2_100"
REPORT_DIRECTORY = PROJECT_ROOT / "reports"

EXCEL_FILENAME = "v2_100_complete_results.xlsx"
CSV_FILENAME = "v2_100_video_results.csv"


# ==================================================
# METRIC CALCULATIONS
# ==================================================

def calculate_metrics(results):
    tp = 0
    tn = 0
    fp = 0
    fn = 0

    for result in results:
        actual = str(result["ground_truth"]).lower()
        predicted = str(result["prediction"]).lower()

        if actual == "accident" and predicted == "accident":
            tp += 1

        elif actual == "normal" and predicted == "normal":
            tn += 1

        elif actual == "normal" and predicted == "accident":
            fp += 1

        elif actual == "accident" and predicted == "normal":
            fn += 1

    total = tp + tn + fp + fn

    accuracy = (
        (tp + tn) / total * 100
        if total > 0 else 0
    )

    precision = (
        tp / (tp + fp) * 100
        if (tp + fp) > 0 else 0
    )

    recall = (
        tp / (tp + fn) * 100
        if (tp + fn) > 0 else 0
    )

    f1 = (
        2 * precision * recall / (precision + recall)
        if (precision + recall) > 0 else 0
    )

    specificity = (
        tn / (tn + fp) * 100
        if (tn + fp) > 0 else 0
    )

    return {
        "TP": tp,
        "TN": tn,
        "FP": fp,
        "FN": fn,
        "Total": total,
        "Accuracy": accuracy,
        "Precision": precision,
        "Recall": recall,
        "F1": f1,
        "Specificity": specificity,
    }


# ==================================================
# READ LOGS
# ==================================================

def read_logs():
    if not LOG_DIRECTORY.exists():
        raise RuntimeError(
            f"Log directory does not exist:\n{LOG_DIRECTORY}"
        )

    log_files = sorted(LOG_DIRECTORY.glob("*.json"))

    if not log_files:
        raise RuntimeError(
            f"No JSON log files found in:\n{LOG_DIRECTORY}"
        )

    results = []

    print("=" * 70)
    print("AI HUB — REPORT GENERATION ONLY")
    print("=" * 70)
    print()
    print(f"Log directory:")
    print(f"  {LOG_DIRECTORY}")
    print()
    print(f"JSON log files found: {len(log_files)}")
    print()

    for log_file in log_files:
        try:
            with open(log_file, "r", encoding="utf-8") as file:
                result = json.load(file)

            required_fields = [
                "video_name",
                "ground_truth",
                "prediction",
            ]

            for field in required_fields:
                if field not in result:
                    raise ValueError(
                        f"Missing required field '{field}'"
                    )

            results.append(result)

        except Exception as error:
            raise RuntimeError(
                f"Could not read log file:\n{log_file}\n"
                f"Error: {error}"
            )

    return results


# ==================================================
# EXCEL REPORT
# ==================================================

def generate_excel(results, metrics):
    REPORT_DIRECTORY.mkdir(
        parents=True,
        exist_ok=True
    )

    output_path = REPORT_DIRECTORY / EXCEL_FILENAME

    workbook = Workbook()

    # ------------------------------------------------
    # VIDEO RESULTS SHEET
    # ------------------------------------------------

    sheet = workbook.active
    sheet.title = "Video Results"

    headers = [
        "Video",
        "Ground Truth",
        "Prediction",
        "TP",
        "TN",
        "FP",
        "FN",
        "First Accident Timestamp",
        "Frames Processed",
        "Detection Count",
        "Vehicle Detection Count",
        "Suspicious Frame Count",
        "Processing Time (seconds)"
    ]

    sheet.append(headers)

    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    for result in results:

        actual = str(
            result["ground_truth"]
        ).lower()

        predicted = str(
            result["prediction"]
        ).lower()

        tp = 1 if (
            actual == "accident"
            and predicted == "accident"
        ) else 0

        tn = 1 if (
            actual == "normal"
            and predicted == "normal"
        ) else 0

        fp = 1 if (
            actual == "normal"
            and predicted == "accident"
        ) else 0

        fn = 1 if (
            actual == "accident"
            and predicted == "normal"
        ) else 0

        sheet.append([
            result.get("video_name"),
            result.get("ground_truth"),
            result.get("prediction"),
            tp,
            tn,
            fp,
            fn,
            result.get("first_accident_timestamp"),
            result.get("frames_processed"),
            result.get("detection_count"),
            result.get("vehicle_detection_count"),
            result.get("suspicious_frame_count"),
            result.get("processing_time_seconds")
        ])

    total_row = sheet.max_row + 2

    sheet.cell(total_row, 1, "TOTAL")

    sheet.cell(
        total_row,
        4,
        metrics["TP"]
    )

    sheet.cell(
        total_row,
        5,
        metrics["TN"]
    )

    sheet.cell(
        total_row,
        6,
        metrics["FP"]
    )

    sheet.cell(
        total_row,
        7,
        metrics["FN"]
    )

    for cell in sheet[total_row]:
        cell.font = Font(bold=True)

    widths = {
        "A": 18,
        "B": 16,
        "C": 16,
        "D": 8,
        "E": 8,
        "F": 8,
        "G": 8,
        "H": 25,
        "I": 18,
        "J": 18,
        "K": 25,
        "L": 25,
        "M": 25,
    }

    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True
            )

    # ------------------------------------------------
    # COMPLETE CALCULATIONS SHEET
    # ------------------------------------------------

    calc = workbook.create_sheet(
        "Complete Calculations"
    )

    calc["A1"] = "COMPLETE METRIC CALCULATIONS"
    calc["A1"].font = Font(
        bold=True,
        size=14
    )

    calc["A3"] = "Metric"
    calc["B3"] = "Formula"
    calc["C3"] = "Result"

    for cell in calc[3]:
        cell.font = Font(bold=True)

    tp = metrics["TP"]
    tn = metrics["TN"]
    fp = metrics["FP"]
    fn = metrics["FN"]

    total = metrics["Total"]

    calc["A4"] = "True Positives (TP)"
    calc["B4"] = "Accident correctly classified as Accident"
    calc["C4"] = tp

    calc["A5"] = "True Negatives (TN)"
    calc["B5"] = "Normal correctly classified as Normal"
    calc["C5"] = tn

    calc["A6"] = "False Positives (FP)"
    calc["B6"] = "Normal incorrectly classified as Accident"
    calc["C6"] = fp

    calc["A7"] = "False Negatives (FN)"
    calc["B7"] = "Accident incorrectly classified as Normal"
    calc["C7"] = fn

    calc["A9"] = "Total Videos"
    calc["B9"] = "TP + TN + FP + FN"
    calc["C9"] = total

    calc["A11"] = "Accuracy"
    calc["B11"] = "(TP + TN) / (TP + TN + FP + FN)"
    calc["C11"] = metrics["Accuracy"] / 100

    calc["A12"] = "Precision"
    calc["B12"] = "TP / (TP + FP)"
    calc["C12"] = metrics["Precision"] / 100

    calc["A13"] = "Recall"
    calc["B13"] = "TP / (TP + FN)"
    calc["C13"] = metrics["Recall"] / 100

    calc["A14"] = "F1 Score"
    calc["B14"] = "2 × Precision × Recall / (Precision + Recall)"
    calc["C14"] = metrics["F1"] / 100

    calc["A15"] = "Specificity"
    calc["B15"] = "TN / (TN + FP)"
    calc["C15"] = metrics["Specificity"] / 100

    for row in range(11, 16):
        calc.cell(row, 3).number_format = "0.00%"

    calc.column_dimensions["A"].width = 25
    calc.column_dimensions["B"].width = 55
    calc.column_dimensions["C"].width = 18

    for row in calc.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True
            )

    workbook.save(output_path)

    return output_path


# ==================================================
# CSV REPORT
# ==================================================

def generate_csv(results):
    REPORT_DIRECTORY.mkdir(
        parents=True,
        exist_ok=True
    )

    output_path = REPORT_DIRECTORY / CSV_FILENAME

    headers = [
        "Video",
        "Ground Truth",
        "Prediction",
        "TP",
        "TN",
        "FP",
        "FN",
        "First Accident Timestamp",
        "Frames Processed",
        "Detection Count",
        "Vehicle Detection Count",
        "Suspicious Frame Count",
        "Processing Time (seconds)"
    ]

    with open(
        output_path,
        "w",
        newline="",
        encoding="utf-8"
    ) as file:

        writer = csv.writer(file)

        writer.writerow(headers)

        for result in results:

            actual = str(
                result["ground_truth"]
            ).lower()

            predicted = str(
                result["prediction"]
            ).lower()

            tp = int(
                actual == "accident"
                and predicted == "accident"
            )

            tn = int(
                actual == "normal"
                and predicted == "normal"
            )

            fp = int(
                actual == "normal"
                and predicted == "accident"
            )

            fn = int(
                actual == "accident"
                and predicted == "normal"
            )

            writer.writerow([
                result.get("video_name"),
                result.get("ground_truth"),
                result.get("prediction"),
                tp,
                tn,
                fp,
                fn,
                result.get("first_accident_timestamp"),
                result.get("frames_processed"),
                result.get("detection_count"),
                result.get("vehicle_detection_count"),
                result.get("suspicious_frame_count"),
                result.get("processing_time_seconds")
            ])

    return output_path


# ==================================================
# MAIN
# ==================================================

def main():

    results = read_logs()

    metrics = calculate_metrics(results)

    print("=" * 70)
    print("FINAL RESULTS")
    print("=" * 70)

    print(f"Videos processed successfully: {len(results)}")
    print()

    print(f"TP: {metrics['TP']}")
    print(f"TN: {metrics['TN']}")
    print(f"FP: {metrics['FP']}")
    print(f"FN: {metrics['FN']}")
    print()

    print(f"Accuracy:    {metrics['Accuracy']:.2f}%")
    print(f"Precision:   {metrics['Precision']:.2f}%")
    print(f"Recall:      {metrics['Recall']:.2f}%")
    print(f"F1 Score:    {metrics['F1']:.2f}%")
    print(f"Specificity: {metrics['Specificity']:.2f}%")
    print()

    excel_path = generate_excel(
        results,
        metrics
    )

    csv_path = generate_csv(
        results
    )

    print("=" * 70)
    print("REPORTS CREATED")
    print("=" * 70)

    print(f"Excel:")
    print(f"  {excel_path}")

    print(f"CSV:")
    print(f"  {csv_path}")

    print()
    print("No videos were processed.")
    print("No detection inference was performed.")
    print("Existing logs were not modified.")


if __name__ == "__main__":
    main()